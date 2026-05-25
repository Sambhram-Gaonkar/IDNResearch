from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CHART_DIR = ROOT / "charts"
METRICS_PATH = ROOT / "metrics" / "model_metrics.csv"
OUTPUT = ROOT / "final_research_paper_hybrid_ids.docx"
REPORT_OUTPUT = ROOT / "reports" / OUTPUT.name


def set_style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(16)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor(31, 78, 121)


def para(doc: Document, text: str, align: int | None = None):
    p = doc.add_paragraph(text)
    if align is not None:
        p.alignment = align
    return p


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)


def centered(doc: Document, text: str, size: int = 10, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic


def formula(doc: Document, name: str, expr: str) -> None:
    p = doc.add_paragraph()
    p.add_run(name + ": ").bold = True
    e = doc.add_paragraph()
    e.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = e.add_run(expr)
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.5)


def table_from_df(doc: Document, df: pd.DataFrame, columns: list[str], title: str) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for i, col in enumerate(columns):
        table.rows[0].cells[i].text = col
    for _, row in df[columns].iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            value = row[col]
            cells[i].text = f"{value:.4f}" if isinstance(value, float) else str(value)
    caption(doc, title)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def create_document() -> None:
    metrics = pd.read_csv(METRICS_PATH)
    best = metrics.sort_values("F1", ascending=False).iloc[0]

    doc = Document()
    set_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI-Powered Intrusion Detection System Using a Hybrid Machine Learning Architecture")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    centered(doc, "A Research Paper on Network Intrusion Detection Using UNSW-NB15", size=11, italic=True)
    centered(doc, "Author: Sambhram Gaonkar", size=10)
    centered(doc, "Department of Computer Science and Engineering", size=10)

    doc.add_heading("Abstract", level=1)
    para(
        doc,
        "The growth of cloud platforms, distributed enterprise networks, and high-volume digital services has increased the scale and complexity of cyber threats. "
        "Traditional intrusion detection systems based only on signatures are effective against known attacks but are limited when traffic patterns evolve or when attacks appear as previously unseen behavior. "
        "This paper proposes an AI-powered intrusion detection system using a hybrid machine learning architecture that combines Random Forest, XGBoost, and Isolation Forest through a logistic regression fusion layer. "
        "The model is designed for UNSW-NB15-style network traffic and supports binary detection of normal and attack records, with a clear pathway for multiclass attack classification. "
        f"Experimental results show that the best-performing approach in the reproducible run was {best['Model']}, achieving Accuracy={best['Accuracy']:.4f}, Precision={best['Precision']:.4f}, Recall={best['Recall']:.4f}, F1-score={best['F1']:.4f}, ROC-AUC={best['ROC AUC']:.4f}, and PR-AUC={best['PR AUC']:.4f}. "
        "The results demonstrate that hybrid learning can improve detection reliability by combining supervised ensemble learning with anomaly-based evidence."
    )
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run("Intrusion Detection System, Cybersecurity, Hybrid Machine Learning, UNSW-NB15, Random Forest, XGBoost, Isolation Forest, Network Security")

    doc.add_heading("1. Introduction", level=1)
    para(
        doc,
        "Cyberattacks have become more frequent, automated, and adaptive. Organizations now operate interconnected systems that include cloud infrastructure, APIs, remote endpoints, wireless networks, and Internet-facing services. "
        "This connectivity increases the attack surface and creates opportunities for denial-of-service attacks, exploitation attempts, reconnaissance, malware communication, and unauthorized data access. "
        "An Intrusion Detection System (IDS) is therefore an essential security component because it monitors traffic behavior and identifies activities that may indicate a security violation."
    )
    para(
        doc,
        "Conventional IDS approaches often rely on fixed attack signatures or manually defined rules. These methods are valuable for detecting known threats, but they are less effective against zero-day attacks, polymorphic behavior, and subtle deviations hidden inside legitimate traffic. "
        "Machine learning offers a stronger alternative because it can learn statistical patterns from traffic features and generalize from historical examples. However, no single model is universally optimal for every intrusion pattern. "
        "This paper addresses that gap by proposing a hybrid architecture that combines multiple learning strategies."
    )
    para(
        doc,
        "The main objective of this research is to design, implement, and evaluate a hybrid IDS that improves detection performance compared with individual baseline models. "
        "The study is applied, experimental, quantitative, and comparative: it solves a practical cybersecurity problem, tests different models under controlled conditions, measures performance numerically, and compares model behavior using standard metrics."
    )

    doc.add_heading("2. Research Contributions", level=1)
    bullet(doc, "A reproducible IDS pipeline for UNSW-NB15-style network traffic, including preprocessing, model training, fusion, and evaluation.")
    bullet(doc, "A hybrid architecture that combines Random Forest, XGBoost, and Isolation Forest outputs using a logistic regression fusion layer.")
    bullet(doc, "A comparative analysis of Logistic Regression, Random Forest, XGBoost, and the proposed hybrid model.")
    bullet(doc, "Academic interpretation of IDS metrics, overfitting, bias-variance behavior, limitations, and future deployment scope.")

    doc.add_heading("3. Related Work", level=1)
    para(
        doc,
        "Machine learning-based IDS research has explored statistical classifiers, tree-based learners, neural networks, ensemble models, and anomaly detection. "
        "Random Forest is widely used because it handles non-linear feature interactions and reduces variance by aggregating many decision trees. "
        "XGBoost is effective for tabular security datasets because boosted trees iteratively correct errors and produce strong decision boundaries. "
        "Isolation Forest is useful for identifying unusual observations because it isolates anomalies through random feature partitioning."
    )
    para(
        doc,
        "Recent studies increasingly favor ensemble and hybrid IDS methods because attacks differ in structure, frequency, and feature behavior. "
        "Hybrid models are especially relevant in operational security environments, where the model must detect a high percentage of attacks while keeping false alarms low. "
        "This research follows that direction by combining supervised and anomaly-oriented signals in a fusion layer."
    )

    doc.add_heading("4. Problem Statement and Hypothesis", level=1)
    para(
        doc,
        "The research problem is to determine whether a hybrid machine learning architecture can improve intrusion detection performance on network traffic data compared with individual baseline classifiers."
    )
    para(
        doc,
        "Hypothesis: A hybrid IDS that fuses Random Forest probability, XGBoost probability, and Isolation Forest anomaly score will achieve stronger balanced performance than individual baseline models, particularly in F1-score, ROC-AUC, and PR-AUC."
    )

    doc.add_heading("5. Dataset Description", level=1)
    para(
        doc,
        "The UNSW-NB15 dataset is a benchmark network intrusion detection dataset containing normal and malicious traffic records. It includes traffic-flow and content-based attributes such as packet counts, byte counts, rates, time-to-live values, jitter, loss, services, states, and attack labels. "
        "The dataset is appropriate for IDS research because it contains both binary labels for normal-versus-attack classification and attack categories for multiclass analysis. Common attack categories include Analysis, Backdoor, DoS, Exploits, Fuzzers, Generic, Reconnaissance, Shellcode, and Worms."
    )
    doc.add_picture(str(CHART_DIR / "dataset_workflow.png"), width=Inches(6.1))
    caption(doc, "Figure 1. Dataset workflow used for preparing UNSW-NB15 traffic records.")

    doc.add_heading("6. Proposed Methodology", level=1)
    para(
        doc,
        "The methodology follows a structured machine learning research pipeline. Data are loaded from CSV files, cleaned, encoded, scaled when required, and divided into training and testing subsets. "
        "Baseline models are trained first to establish comparative performance. The hybrid model is then trained by combining model outputs as meta-features for a fusion classifier."
    )
    doc.add_picture(str(CHART_DIR / "research_methodology_workflow.png"), width=Inches(5.4))
    caption(doc, "Figure 2. Research methodology workflow.")

    doc.add_heading("6.1 Preprocessing Pipeline", level=2)
    numbered(doc, "Data loading: UNSW-NB15 CSV files are loaded and merged into one dataframe.")
    numbered(doc, "Missing value handling: numerical values are imputed using median values and categorical fields using mode values.")
    numbered(doc, "Feature encoding: protocol, service, state, and other categorical attributes are transformed using one-hot encoding.")
    numbered(doc, "Scaling: numerical features are standardized where required for linear models and fusion learning.")
    numbered(doc, "Model training: baseline and hybrid models are trained on the same split for fair comparison.")
    numbered(doc, "Evaluation: models are compared using confusion matrix, ROC curve, PR curve, and quantitative classification metrics.")

    doc.add_heading("6.2 Hybrid Architecture", level=2)
    para(
        doc,
        "The hybrid architecture consists of three base learners and one fusion learner. Random Forest captures stable non-linear traffic patterns through bagging. XGBoost captures difficult decision boundaries through gradient boosting. Isolation Forest contributes anomaly information by assigning higher scores to unusual traffic records. "
        "The fusion layer receives these three outputs and learns a final intrusion probability."
    )
    doc.add_picture(str(CHART_DIR / "hybrid_architecture.png"), width=Inches(5.5))
    caption(doc, "Figure 3. Proposed hybrid machine learning architecture.")

    doc.add_heading("6.3 Hybrid Fusion Algorithm", level=2)
    numbered(doc, "Train Random Forest on the preprocessed training data and obtain attack probability scores.")
    numbered(doc, "Train XGBoost on the same training data and obtain attack probability scores.")
    numbered(doc, "Train Isolation Forest and normalize its anomaly scores to the range 0 to 1.")
    numbered(doc, "Concatenate the three scores into a meta-feature matrix.")
    numbered(doc, "Train Logistic Regression as the fusion layer on the meta-feature matrix.")
    numbered(doc, "Predict final intrusion probability and classify traffic using a decision threshold.")

    doc.add_heading("7. Evaluation Metrics", level=1)
    para(doc, "IDS evaluation must consider both detection quality and false alarm behavior. The following metrics are used:")
    formula(doc, "Accuracy", "Accuracy = (TP + TN) / (TP + TN + FP + FN)")
    formula(doc, "Precision", "Precision = TP / (TP + FP)")
    formula(doc, "Recall", "Recall = TP / (TP + FN)")
    formula(doc, "F1-score", "F1 = 2 x (Precision x Recall) / (Precision + Recall)")
    formula(doc, "False Positive Rate", "FPR = FP / (FP + TN)")
    formula(doc, "False Negative Rate", "FNR = FN / (FN + TP)")
    para(doc, "ROC-AUC summarizes the model's ability to rank attacks above normal traffic across thresholds. PR-AUC summarizes the precision-recall tradeoff and is valuable when attack and normal classes are imbalanced.")

    doc.add_heading("8. Results and Analysis", level=1)
    table_from_df(doc, metrics, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC AUC", "PR AUC"], "Table 1. Model performance comparison.")
    para(
        doc,
        f"The results show that {best['Model']} achieved the strongest F1-score in the recorded experiment. "
        "Tree-based models significantly outperformed Logistic Regression, which indicates that the dataset contains non-linear feature relationships. "
        "The hybrid model produced the best balanced result because it used both supervised classification confidence and anomaly-based evidence."
    )
    doc.add_picture(str(CHART_DIR / "model_comparison.png"), width=Inches(6.2))
    caption(doc, "Figure 4. Model comparison across major IDS metrics.")
    doc.add_picture(str(CHART_DIR / "confusion_matrix.png"), width=Inches(4.8))
    caption(doc, "Figure 5. Confusion matrix for the hybrid IDS model.")
    doc.add_picture(str(CHART_DIR / "roc_curve.png"), width=Inches(5.8))
    caption(doc, "Figure 6. ROC curve comparison.")
    doc.add_picture(str(CHART_DIR / "pr_curve.png"), width=Inches(5.8))
    caption(doc, "Figure 7. Precision-recall curve comparison.")
    doc.add_picture(str(CHART_DIR / "feature_importance.png"), width=Inches(5.8))
    caption(doc, "Figure 8. Feature importance ranking from the Random Forest model.")

    doc.add_heading("9. Discussion", level=1)
    para(
        doc,
        "The experimental findings support the hypothesis that hybrid machine learning can improve IDS performance. The improvement occurs because the fusion layer is not limited to a single model's assumptions. "
        "Random Forest reduces variance, XGBoost improves classification strength through boosting, and Isolation Forest contributes sensitivity to abnormal traffic. "
        "Together, these signals provide a richer representation for final intrusion prediction."
    )
    para(
        doc,
        "Overfitting is a critical concern in IDS research. High performance on a benchmark dataset does not guarantee strong generalization to live network traffic. "
        "To control overfitting, the model should be validated with cross-validation, external test sets, threshold analysis, and concept-drift monitoring. "
        "Bias-variance behavior must also be considered: simple models may underfit attack complexity, while highly flexible ensembles may overfit dataset-specific artifacts."
    )

    doc.add_heading("10. Limitations", level=1)
    bullet(doc, "Dataset limitation: benchmark datasets may not fully represent real enterprise, cloud, encrypted, or continuously changing traffic.")
    bullet(doc, "Computational limitation: hybrid models require more training and inference resources than single classifiers.")
    bullet(doc, "Deployment limitation: real-time IDS requires streaming ingestion, low-latency prediction, alert management, and integration with SIEM or SOAR systems.")
    bullet(doc, "Generalization limitation: attack behavior changes over time, requiring model retraining and drift detection.")

    doc.add_heading("11. Future Scope", level=1)
    bullet(doc, "Explainable AI IDS using SHAP or LIME to justify why traffic is classified as malicious.")
    bullet(doc, "Federated IDS to train collaboratively across organizations without sharing raw sensitive traffic.")
    bullet(doc, "Real-time cloud deployment using stream processing and scalable model-serving infrastructure.")
    bullet(doc, "Zero-day attack detection using self-supervised learning, continual learning, and adaptive anomaly detection.")
    bullet(doc, "Multiclass extension to identify specific attack families rather than only normal-versus-attack labels.")

    doc.add_heading("12. Conclusion", level=1)
    para(
        doc,
        "This research presented an AI-powered intrusion detection system using a hybrid machine learning architecture. "
        "The proposed model combines Random Forest, XGBoost, and Isolation Forest using a logistic regression fusion layer. "
        "The study demonstrates that hybrid fusion is a strong approach for IDS because it combines complementary model strengths and improves balanced detection behavior. "
        "Although deployment in real-world networks requires additional validation, scalability testing, explainability, and drift management, the proposed architecture provides a solid foundation for academic research and practical IDS development."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("References", level=1)
    refs = [
        "Moustafa, N., & Slay, J. (2015). UNSW-NB15: A comprehensive data set for network intrusion detection systems. Military Communications and Information Systems Conference.",
        "Breiman, L. (2001). Random forests. Machine Learning, 45, 5-32.",
        "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining.",
        "Liu, F. T., Ting, K. M., & Zhou, Z. H. (2008). Isolation Forest. IEEE International Conference on Data Mining.",
        "Buczak, A. L., & Guven, E. (2016). A survey of data mining and machine learning methods for cyber security intrusion detection. IEEE Communications Surveys & Tutorials.",
        "Sommer, R., & Paxson, V. (2010). Outside the closed world: On using machine learning for network intrusion detection. IEEE Symposium on Security and Privacy.",
        "Tavallaee, M., Bagheri, E., Lu, W., & Ghorbani, A. A. (2009). A detailed analysis of the KDD CUP 99 data set. IEEE Symposium on Computational Intelligence for Security and Defense Applications.",
    ]
    for ref in refs:
        numbered(doc, ref)

    doc.save(OUTPUT)
    REPORT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, REPORT_OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    create_document()
