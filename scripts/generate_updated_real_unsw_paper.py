from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".matplotlib-cache"))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import IsolationForest, RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    ConfusionMatrixDisplay,
    accuracy_score,
    average_precision_score,
    confusion_matrix,
    f1_score,
    precision_recall_curve,
    precision_score,
    recall_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from xgboost import XGBClassifier


DATA_DIR = ROOT / "data"
CHART_DIR = ROOT / "charts"
METRIC_DIR = ROOT / "metrics"
OUTPUT = ROOT / "updated_hybrid_ids_research_paper_real_unsw_2021_2026.docx"
REPORT_OUTPUT = ROOT / "reports" / OUTPUT.name

TRAIN_PATH = DATA_DIR / "UNSW_NB15_training-set.csv"
TEST_PATH = DATA_DIR / "UNSW_NB15_testing-set.csv"
METRICS_CSV = METRIC_DIR / "real_unsw_model_metrics.csv"
METRICS_JSON = METRIC_DIR / "real_unsw_model_metrics.json"


REFERENCES = [
    "Mabulage, D., & Athuraliya, B. (2026). LiteShield: Hybrid Feature Selection-Driven Lightweight Intrusion Detection for Resource-Constrained IoT Networks. arXiv:2605.02987. https://arxiv.org/abs/2605.02987",
    "Hossain, M. Z., et al. (2026). Assessing Generalisation Capability of Machine Learning Models for Intrusion Detection. arXiv:2605.04407. https://arxiv.org/abs/2605.04407",
    "Alhusseini, M. M., Rouhi, A., & Feizi-Derakhshi, M.-R. (2026). AI-Powered Hybrid Intrusion Detection Framework for Cloud Security Using Novel Metaheuristic Optimization. arXiv:2601.01134. https://arxiv.org/abs/2601.01134",
    "Hussein, S. A., & Repas, S. R. (2026). A Hybrid Intrusion Detection Framework Using Deep Autoencoder and Machine Learning Models. AI, 7(2), 39. https://www.mdpi.com/2673-2688/7/2/39",
    "A Rigorous Comparative Study of Supervised Machine Learning Techniques for Network Anomaly Detection: Empirical Insights from the UNSW-NB15 Dataset. (2026). Computers, 15(5), 285. https://www.mdpi.com/2073-431X/15/5/285",
    "Farhan, M., et al. (2025). Network-based intrusion detection using deep learning technique. Scientific Reports, 15, 25550. https://www.nature.com/articles/s41598-025-08770-0",
    "Ali, M., et al. (2025). Botnet detection in internet of things using stacked ensemble learning model. Scientific Reports, 15, 21012. https://www.nature.com/articles/s41598-025-02008-9",
    "Li, Z., et al. (2025). Generating detectors from anomaly samples via negative selection for network intrusion detection. Scientific Reports, 15, 36456. https://www.nature.com/articles/s41598-025-20516-6",
    "xIDS-EnsembleGuard: An Explainable Ensemble Learning-based Intrusion Detection System. (2025). arXiv:2503.00615. https://arxiv.org/abs/2503.00615",
    "Network Intrusion Detection through Stacked Machine Learning Models on UNSW-NB15 Data Set. (2025). SCITEPRESS. https://www.scitepress.org/Papers/2025/139326/139326.pdf",
    "A new intrusion detection method using ensemble classification and feature selection. (2025). Scientific Reports. https://www.nature.com/articles/s41598-025-98604-w",
    "Varzaneh, Z. A., & Hosseini, S. (2024). An improved equilibrium optimization algorithm for feature selection problem in network intrusion detection. Scientific Reports, 14, 18696. https://www.nature.com/articles/s41598-024-67488-7",
    "Explainable AI for Comparative Analysis of Intrusion Detection Models. (2024). arXiv:2406.09684. https://arxiv.org/abs/2406.09684",
    "Zoghi, Z., & Serpen, G. (2022). Ensemble Classifier Design Tuned to Dataset Characteristics for Network Intrusion Detection. arXiv:2205.06177. https://arxiv.org/abs/2205.06177",
    "Yin, Y., Jang-Jaccard, J., Xu, W., Singh, A., Zhu, J., Sabrina, F., & Kwak, J. (2022). IGRF-RFE: A Hybrid Feature Selection Method for MLP-based Network Intrusion Detection on UNSW-NB15 Dataset. arXiv:2203.16365. https://arxiv.org/abs/2203.16365",
    "UNSW-NB15 Computer Security Dataset: Analysis through Visualization. (2021). arXiv:2101.05067. https://arxiv.org/abs/2101.05067",
]


def set_style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(16)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor(31, 78, 121)


def para(doc: Document, text: str):
    return doc.add_paragraph(text)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def centered(doc: Document, text: str, size: int = 10, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def formula(doc: Document, name: str, expr: str) -> None:
    p = doc.add_paragraph()
    p.add_run(name + ": ").bold = True
    e = doc.add_paragraph()
    e.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = e.add_run(expr)
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.5)


def encoder() -> OneHotEncoder:
    try:
        return OneHotEncoder(handle_unknown="ignore", sparse_output=True)
    except TypeError:
        return OneHotEncoder(handle_unknown="ignore")


def load_real_split() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not TRAIN_PATH.exists() or not TEST_PATH.exists():
        raise FileNotFoundError(
            "Missing real UNSW-NB15 files. Expected data/UNSW_NB15_training-set.csv "
            "and data/UNSW_NB15_testing-set.csv."
        )
    train = pd.read_csv(TRAIN_PATH)
    test = pd.read_csv(TEST_PATH)
    return train, test


def build_preprocessor(df: pd.DataFrame) -> ColumnTransformer:
    feature_df = df.drop(columns=["label", "attack_cat", "id"], errors="ignore")
    numeric = feature_df.select_dtypes(include=["number", "bool"]).columns.tolist()
    categorical = [col for col in feature_df.columns if col not in numeric]
    numeric_pipe = Pipeline([("imputer", SimpleImputer(strategy="median")), ("scaler", StandardScaler())])
    categorical_pipe = Pipeline([("imputer", SimpleImputer(strategy="most_frequent")), ("encoder", encoder())])
    return ColumnTransformer([("num", numeric_pipe, numeric), ("cat", categorical_pipe, categorical)])


def normalized_anomaly_scores(model: IsolationForest, x) -> np.ndarray:
    scores = -model.decision_function(x)
    return (scores - scores.min()) / max(scores.max() - scores.min(), 1e-9)


def run_analysis() -> tuple[pd.DataFrame, dict[str, np.ndarray], np.ndarray, list[str], pd.DataFrame, pd.DataFrame]:
    train, test = load_real_split()
    y_train = train["label"].astype(int)
    y_test = test["label"].astype(int)
    x_train = train.drop(columns=["label", "id"], errors="ignore")
    x_test = test.drop(columns=["label", "id"], errors="ignore")

    preprocessor = build_preprocessor(train)
    x_train_t = preprocessor.fit_transform(x_train)
    x_test_t = preprocessor.transform(x_test)

    scale_pos_weight = max((y_train == 0).sum() / max((y_train == 1).sum(), 1), 1e-9)
    models = {
        "Logistic Regression": LogisticRegression(max_iter=600, class_weight="balanced", random_state=42),
        "Random Forest": RandomForestClassifier(
            n_estimators=160, class_weight="balanced_subsample", random_state=42, n_jobs=1
        ),
        "XGBoost": XGBClassifier(
            n_estimators=180,
            max_depth=6,
            learning_rate=0.08,
            subsample=0.9,
            colsample_bytree=0.9,
            eval_metric="logloss",
            random_state=42,
            n_jobs=2,
            scale_pos_weight=scale_pos_weight,
        ),
    }

    probabilities: dict[str, np.ndarray] = {}
    train_probabilities: dict[str, np.ndarray] = {}
    for name, model in models.items():
        model.fit(x_train_t, y_train)
        probabilities[name] = model.predict_proba(x_test_t)[:, 1]
        train_probabilities[name] = model.predict_proba(x_train_t)[:, 1]

    iso = IsolationForest(n_estimators=120, contamination="auto", random_state=42, n_jobs=1)
    iso.fit(x_train_t)
    iso_train = normalized_anomaly_scores(iso, x_train_t)
    iso_test = normalized_anomaly_scores(iso, x_test_t)

    stacked_train = np.column_stack([train_probabilities["Random Forest"], train_probabilities["XGBoost"], iso_train])
    stacked_test = np.column_stack([probabilities["Random Forest"], probabilities["XGBoost"], iso_test])
    fusion = LogisticRegression(max_iter=600, random_state=42)
    fusion.fit(stacked_train, y_train)
    probabilities["Hybrid RF + XGBoost + Isolation Forest"] = fusion.predict_proba(stacked_test)[:, 1]

    rows = []
    for name, proba in probabilities.items():
        pred = (proba >= 0.5).astype(int)
        rows.append(
            {
                "Model": name,
                "Accuracy": accuracy_score(y_test, pred),
                "Precision": precision_score(y_test, pred, zero_division=0),
                "Recall": recall_score(y_test, pred, zero_division=0),
                "F1": f1_score(y_test, pred, zero_division=0),
                "ROC AUC": roc_auc_score(y_test, proba),
                "PR AUC": average_precision_score(y_test, proba),
                "False Positive Rate": confusion_matrix(y_test, pred)[0, 1] / max((y_test == 0).sum(), 1),
                "False Negative Rate": confusion_matrix(y_test, pred)[1, 0] / max((y_test == 1).sum(), 1),
            }
        )
    metrics = pd.DataFrame(rows).sort_values("F1", ascending=False)
    METRIC_DIR.mkdir(exist_ok=True)
    metrics.to_csv(METRICS_CSV, index=False)
    METRICS_JSON.write_text(json.dumps(metrics.to_dict(orient="records"), indent=2), encoding="utf-8")
    feature_names = list(preprocessor.get_feature_names_out())
    return metrics, probabilities, y_test.to_numpy(), feature_names, train, test


def save_fig(path: Path) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()


def make_charts(metrics: pd.DataFrame, probabilities: dict[str, np.ndarray], y_test: np.ndarray, feature_names: list[str]) -> None:
    CHART_DIR.mkdir(exist_ok=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "axes.titlesize": 14,
        "axes.labelsize": 10,
        "xtick.labelsize": 9,
        "ytick.labelsize": 9,
    })
    palette = ["#2f6f8f", "#4f9d69", "#d9923b", "#6b5b95"]
    model_labels = {
        "Hybrid RF + XGBoost + Isolation Forest": "Hybrid RF + XGBoost\n+ Isolation Forest",
        "Logistic Regression": "Logistic\nRegression",
    }

    plot_df = metrics.set_index("Model")[["Accuracy", "Precision", "Recall", "F1", "ROC AUC", "PR AUC"]]
    fig, ax = plt.subplots(figsize=(8.6, 3.9))
    image = ax.imshow(plot_df.values, cmap="YlGnBu", vmin=0.78, vmax=1.0, aspect="auto")
    ax.set_xticks(range(len(plot_df.columns)), plot_df.columns)
    ax.set_yticks(range(len(plot_df.index)), [model_labels.get(label, label) for label in plot_df.index])
    ax.set_title("Model Performance on Real UNSW-NB15")
    for row_idx in range(plot_df.shape[0]):
        for col_idx in range(plot_df.shape[1]):
            value = plot_df.iloc[row_idx, col_idx]
            text_color = "white" if value >= 0.965 else "#102a43"
            ax.text(col_idx, row_idx, f"{value:.3f}", ha="center", va="center", color=text_color, fontsize=8.5)
    for spine in ax.spines.values():
        spine.set_visible(False)
    cbar = fig.colorbar(image, ax=ax, fraction=0.035, pad=0.03)
    cbar.set_label("Score", rotation=270, labelpad=12)
    save_fig(CHART_DIR / "real_unsw_model_comparison.png")

    fig, ax = plt.subplots(figsize=(5.4, 4.5))
    best_name = metrics.iloc[0]["Model"]
    best = metrics.iloc[0]
    normal_count = int((y_test == 0).sum())
    attack_count = int((y_test == 1).sum())
    fp = int(round(best["False Positive Rate"] * normal_count))
    fn = int(round(best["False Negative Rate"] * attack_count))
    tn = normal_count - fp
    tp = attack_count - fn
    cm = np.array([[tn, fp], [fn, tp]])
    display = ConfusionMatrixDisplay(confusion_matrix=cm, display_labels=["Normal", "Attack"])
    display.plot(cmap="Blues", colorbar=False, ax=ax, values_format=",d")
    ax.set_title(f"Confusion Matrix: {best_name}")
    ax.set_xlabel("Predicted label")
    ax.set_ylabel("True label")
    save_fig(CHART_DIR / "real_unsw_confusion_matrix.png")

    auc_df = metrics.sort_values("ROC AUC", ascending=True)
    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    ax.barh([model_labels.get(label, label) for label in auc_df["Model"]], auc_df["ROC AUC"], color=palette)
    ax.set_xlim(0.94, 1.0)
    ax.set_xlabel("ROC-AUC")
    ax.set_title("ROC-AUC Comparison")
    ax.grid(axis="x", alpha=0.2)
    for idx, value in enumerate(auc_df["ROC AUC"]):
        ax.text(value + 0.001, idx, f"{value:.3f}", va="center", fontsize=9)
    ax.spines[["top", "right", "left"]].set_visible(False)
    save_fig(CHART_DIR / "real_unsw_roc_curve.png")

    pr_df = metrics.sort_values("PR AUC", ascending=True)
    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    ax.barh([model_labels.get(label, label) for label in pr_df["Model"]], pr_df["PR AUC"], color=palette)
    ax.set_xlim(0.96, 1.0)
    ax.set_xlabel("PR-AUC")
    ax.set_title("Precision-Recall AUC Comparison")
    ax.grid(axis="x", alpha=0.2)
    for idx, value in enumerate(pr_df["PR AUC"]):
        ax.text(value + 0.0008, idx, f"{value:.3f}", va="center", fontsize=9)
    ax.spines[["top", "right", "left"]].set_visible(False)
    save_fig(CHART_DIR / "real_unsw_pr_curve.png")

    fig, ax = plt.subplots(figsize=(8.8, 4.6))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    boxes = [
        ("Real UNSW-NB15\ntrain/test CSVs", 0.14, 0.72),
        ("Preprocessing\nimpute, encode, scale", 0.38, 0.72),
        ("Base learners\nRF, XGBoost, IF", 0.62, 0.72),
        ("Fusion layer\nlogistic regression", 0.62, 0.35),
        ("IDS decision\nNormal / Attack", 0.86, 0.35),
    ]
    for label, x, y in boxes:
        ax.text(
            x,
            y,
            label,
            ha="center",
            va="center",
            fontsize=10,
            linespacing=1.25,
            bbox=dict(boxstyle="round,pad=0.55,rounding_size=0.08", fc="#f7fbff", ec="#2f6f8f", lw=1.5),
        )
    arrows = [
        ((0.24, 0.72), (0.30, 0.72)),
        ((0.48, 0.72), (0.54, 0.72)),
        ((0.62, 0.62), (0.62, 0.45)),
        ((0.72, 0.35), (0.78, 0.35)),
    ]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", lw=1.8, color="#3d4852"))
    ax.text(0.62, 0.54, "probabilities + anomaly score", ha="center", va="center", fontsize=8.5, color="#52616b")
    ax.set_title("Hybrid IDS Analysis Workflow", fontsize=14, pad=10)
    save_fig(CHART_DIR / "real_unsw_research_workflow.png")


def table_from_df(doc: Document, df: pd.DataFrame, columns: list[str], title: str) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for i, col in enumerate(columns):
        table.rows[0].cells[i].text = col
    for _, row in df[columns].iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            value = row[col]
            cells[i].text = f"{value:.4f}" if isinstance(value, float) else str(value)
    caption(doc, title)


def add_dataset_table(doc: Document, train: pd.DataFrame, test: pd.DataFrame) -> None:
    rows = [
        ("Training split", len(train), int(train["label"].eq(0).sum()), int(train["label"].eq(1).sum())),
        ("Testing split", len(test), int(test["label"].eq(0).sum()), int(test["label"].eq(1).sum())),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for idx, col in enumerate(["Split", "Records", "Normal", "Attack"]):
        table.rows[0].cells[idx].text = col
    for split, records, normal, attack in rows:
        cells = table.add_row().cells
        cells[0].text = split
        cells[1].text = f"{records:,}"
        cells[2].text = f"{normal:,}"
        cells[3].text = f"{attack:,}"
    caption(doc, "Table 1. Real UNSW-NB15 split used for the updated analysis.")


def create_document(metrics: pd.DataFrame, train: pd.DataFrame, test: pd.DataFrame) -> None:
    best = metrics.iloc[0]
    doc = Document()
    set_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI-Powered Hybrid Intrusion Detection System on the Real UNSW-NB15 Dataset")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    centered(doc, "Updated research paper with 2021-2026 literature and real-dataset evaluation", size=11, italic=True)
    centered(doc, "Author: Shivani Nayak", size=10)
    centered(doc, "Department of Computer Science and Engineering", size=10)

    doc.add_heading("Abstract", level=1)
    para(doc, "Modern enterprise, cloud, and IoT networks generate large volumes of traffic where malicious behavior can be hidden inside legitimate activity. Signature-based intrusion detection systems are useful for known attacks, but they are limited when attack behavior changes, when traffic classes are imbalanced, or when previously unseen behavior appears. This updated study develops a hybrid machine learning intrusion detection system using the real UNSW-NB15 training and testing CSV files rather than the earlier synthetic sample dataset.")
    para(doc, f"The proposed architecture combines Random Forest, XGBoost, and Isolation Forest signals through a logistic regression fusion layer. On the real UNSW-NB15 test split, the strongest model in this run was {best['Model']}, with Accuracy={best['Accuracy']:.4f}, Precision={best['Precision']:.4f}, Recall={best['Recall']:.4f}, F1-score={best['F1']:.4f}, ROC-AUC={best['ROC AUC']:.4f}, and PR-AUC={best['PR AUC']:.4f}. The findings support the use of hybrid supervised and anomaly-aware learning for stronger IDS decision-making, while also showing why benchmark results must be interpreted with attention to dataset split, class imbalance, and operational false-alarm cost.")
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run("Intrusion Detection System, UNSW-NB15, Hybrid Machine Learning, Random Forest, XGBoost, Isolation Forest, Explainable AI, Cybersecurity")

    doc.add_heading("1. Introduction", level=1)
    para(doc, "Network intrusion detection remains important because organizations now operate across cloud services, remote endpoints, APIs, IoT devices, and software-defined infrastructure. These environments increase the number of exposed services and produce traffic that changes rapidly over time. Attackers can use reconnaissance, exploit attempts, denial-of-service traffic, malware communication, and credential abuse in ways that do not always match fixed signatures.")
    para(doc, "Machine learning helps by learning traffic patterns from labeled records, but individual models have different weaknesses. Linear models are interpretable but may miss non-linear feature interactions. Bagging methods such as Random Forest are stable but can underfit rare attack boundaries. Boosted trees such as XGBoost often produce strong tabular performance but can overfit dataset artifacts. Anomaly detectors can identify unusual behavior, but they do not always separate benign novelty from actual attacks. This research therefore uses a hybrid design so the final IDS decision is based on complementary signals rather than one model family.")

    doc.add_heading("2. Research Objectives and Contributions", level=1)
    bullet(doc, "Replace the earlier sample-dataset evaluation with the real UNSW-NB15 training and testing CSV files.")
    bullet(doc, "Use recent IDS literature from 2021 to 2026 as the main bibliography instead of relying on older references.")
    bullet(doc, "Compare Logistic Regression, Random Forest, XGBoost, and a hybrid RF + XGBoost + Isolation Forest fusion model.")
    bullet(doc, "Improve the research explanation with clearer dataset provenance, preprocessing, hybrid fusion logic, evaluation interpretation, and deployment limitations.")
    bullet(doc, "Include cleaner figures generated from the updated real-data experiment.")

    doc.add_heading("3. Recent Literature Review: 2021-2026", level=1)
    para(doc, "Recent intrusion detection research has shifted from single benchmark classifiers toward hybrid, explainable, lightweight, and deployment-aware systems. The 2021 UNSW-NB15 visualization study reinforced the importance of understanding feature distributions before modeling. Zoghi and Serpen (2022) showed that classifier design should be tuned to dataset characteristics, while Yin et al. (2022) demonstrated that hybrid feature selection can reduce feature dimensionality and improve multiclass detection.")
    para(doc, "From 2024 onward, explainability and robustness became more central. Comparative XAI work on UNSW-NB15 connected model accuracy with interpretable evidence, while equilibrium-optimization research treated feature selection as a core IDS performance factor. The 2025 studies broadened the field through deep learning, stacked ensembles, negative selection, botnet detection, and ensemble feature-selection methods. These papers repeatedly show that accuracy alone is not enough; IDS models must also control false positives, false negatives, minority attack behavior, and operational cost.")
    para(doc, "The 2026 papers are directly aligned with this project. LiteShield focuses on lightweight hybrid feature selection for constrained environments, Hossain et al. study generalization limits, Alhusseini et al. propose optimized hybrid IDS for cloud security, and Hussein and Repas combine representation learning with machine learning classifiers. Together, this recent literature supports the design choice used here: combine strong supervised tabular learners with anomaly-oriented evidence and evaluate the result using threshold-sensitive metrics.")

    doc.add_heading("4. Dataset and Real-Data Replacement", level=1)
    para(doc, "The updated experiment uses the real UNSW-NB15 training and testing CSV files: data/UNSW_NB15_training-set.csv and data/UNSW_NB15_testing-set.csv. These files replace the previous sample_unsw_nb15_like.csv demonstration dataset for the new analysis. The sample file remains in the project only as a lightweight fallback for quick execution tests; it is not the basis for the updated research conclusions.")
    para(doc, "UNSW-NB15 contains flow-level traffic attributes and labels for normal traffic plus nine attack categories: Analysis, Backdoor, DoS, Exploits, Fuzzers, Generic, Reconnaissance, Shellcode, and Worms. The binary label is used in this paper for normal-versus-attack classification. The attack_cat field is preserved for future multiclass analysis and attack-family reporting.")
    add_dataset_table(doc, train, test)
    doc.add_picture(str(CHART_DIR / "real_unsw_research_workflow.png"), width=Inches(6.2))
    caption(doc, "Figure 1. Updated research workflow using the real UNSW-NB15 dataset.")

    doc.add_heading("5. Proposed Methodology", level=1)
    para(doc, "The methodology follows an applied quantitative research design. The official training split is used for model fitting and the official testing split is used for evaluation. Numerical features are median-imputed and standardized. Categorical features such as protocol, service, and connection state are mode-imputed and one-hot encoded. The id column is removed because it identifies records rather than network behavior. The attack_cat column is excluded from binary model features to avoid label leakage.")
    para(doc, "Four model groups are evaluated. Logistic Regression is included as a simple linear baseline. Random Forest represents bagged decision-tree learning and is useful for non-linear tabular patterns. XGBoost represents boosted tree learning and is often strong on structured IDS datasets. Isolation Forest is used inside the hybrid model to contribute an anomaly score, which is normalized and combined with supervised probabilities.")
    para(doc, "The hybrid model has two stages. First, Random Forest and XGBoost produce attack probabilities, while Isolation Forest produces anomaly scores. Second, these three outputs are stacked into a compact meta-feature matrix. A logistic regression fusion layer learns how to weight the signals and outputs the final intrusion probability. This design keeps the fusion layer interpretable while allowing the base learners to capture richer traffic behavior.")

    doc.add_heading("6. Evaluation Metrics", level=1)
    para(doc, "Intrusion detection evaluation must consider both attack detection and false alarm behavior. Accuracy can be misleading when normal and attack classes are imbalanced, so this paper also reports precision, recall, F1-score, ROC-AUC, PR-AUC, false positive rate, and false negative rate.")
    formula(doc, "Precision", "Precision = TP / (TP + FP)")
    formula(doc, "Recall", "Recall = TP / (TP + FN)")
    formula(doc, "F1-score", "F1 = 2 x (Precision x Recall) / (Precision + Recall)")
    formula(doc, "False Positive Rate", "FPR = FP / (FP + TN)")
    formula(doc, "False Negative Rate", "FNR = FN / (FN + TP)")
    para(doc, "For IDS use, recall measures how many actual attacks are detected, while precision measures how many generated alerts are truly malicious. FPR is important because high false alarms can overwhelm analysts. FNR is critical because missed attacks may allow compromise. ROC-AUC measures ranking quality across thresholds, and PR-AUC is especially useful when the attack distribution is imbalanced.")

    doc.add_heading("7. Results and Analysis", level=1)
    table_from_df(doc, metrics, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC AUC", "PR AUC", "False Positive Rate", "False Negative Rate"], "Table 2. Model performance on the real UNSW-NB15 test split.")
    para(doc, f"The best F1-score was achieved by {best['Model']}. This result indicates the strongest balance between precision and recall for the default threshold used in the experiment. The comparison also shows whether the hybrid fusion layer improved over the strongest individual base learner or whether one supervised learner already captured most of the useful decision boundary.")
    para(doc, "The confusion matrix explains the operational meaning of the metrics. False positives represent benign traffic that would become analyst workload. False negatives represent attacks that the IDS missed. In practice, the final threshold should be selected according to the organization's risk tolerance; a critical infrastructure environment may accept more false positives to reduce missed attacks, while a high-volume enterprise SOC may need a stricter alert budget.")
    doc.add_picture(str(CHART_DIR / "real_unsw_model_comparison.png"), width=Inches(6.2))
    caption(doc, "Figure 2. Model comparison across major IDS metrics.")
    doc.add_picture(str(CHART_DIR / "real_unsw_confusion_matrix.png"), width=Inches(4.8))
    caption(doc, "Figure 3. Confusion matrix for the strongest model in the real-data run.")
    doc.add_picture(str(CHART_DIR / "real_unsw_roc_curve.png"), width=Inches(5.8))
    caption(doc, "Figure 4. ROC-AUC comparison across evaluated models.")
    doc.add_picture(str(CHART_DIR / "real_unsw_pr_curve.png"), width=Inches(5.8))
    caption(doc, "Figure 5. Precision-recall AUC comparison across evaluated models.")

    doc.add_heading("8. Discussion", level=1)
    para(doc, "The updated results are more defensible than the earlier sample-data results because they are produced from the real UNSW-NB15 split. Real benchmark data contains class overlap, feature noise, categorical protocol behavior, and attack-category imbalance that a small synthetic file cannot represent. This makes the reported performance more meaningful for academic discussion.")
    para(doc, "The hybrid model is theoretically useful because it combines classification confidence with anomaly evidence. If Random Forest and XGBoost agree strongly, the fusion layer can make a high-confidence decision. If supervised probabilities are uncertain but the Isolation Forest score is unusual, the fusion layer can still incorporate anomaly evidence. This is important for intrusion detection because new or rare attacks may not always resemble frequent training examples.")
    para(doc, "However, high benchmark performance should not be interpreted as deployment readiness. Live traffic may differ from UNSW-NB15 because of encryption, routing policy, cloud service behavior, user geography, and concept drift. A production IDS would require periodic retraining, feature monitoring, calibrated thresholds, alert triage integration, and external validation on newer network traces.")

    doc.add_heading("9. Limitations", level=1)
    bullet(doc, "The paper evaluates binary classification only; attack-family multiclass classification remains future work.")
    bullet(doc, "UNSW-NB15 is a benchmark dataset and cannot fully represent every modern cloud, enterprise, or IoT network.")
    bullet(doc, "The fusion threshold is fixed at 0.5; operational deployment should tune the threshold using alert cost and risk tolerance.")
    bullet(doc, "The model does not yet include SHAP/LIME explanations, drift detection, streaming inference, or SIEM integration.")

    doc.add_heading("10. Future Scope", level=1)
    bullet(doc, "Extend the pipeline to multiclass attack_cat prediction for attack-family identification.")
    bullet(doc, "Add SHAP-based explanations for top features behind each malicious classification.")
    bullet(doc, "Evaluate cross-dataset generalization using CIC-IDS2017, BoT-IoT, ToN-IoT, or NF-UNSW-NB15.")
    bullet(doc, "Tune thresholds using cost-sensitive IDS objectives instead of a fixed 0.5 cutoff.")
    bullet(doc, "Develop a streaming prototype that produces alerts with timestamp, predicted class, confidence, and explanation.")

    doc.add_heading("11. Conclusion", level=1)
    para(doc, "This updated research paper presents a hybrid machine learning intrusion detection system evaluated on the real UNSW-NB15 training and testing files. The work replaces the sample-dataset basis with a stronger experimental foundation, updates the literature review to 2021-2026 papers, and explains the modeling pipeline in greater detail. The hybrid approach is a practical research direction because it combines supervised ensemble probabilities with anomaly detection evidence, but final deployment still requires multiclass testing, explainability, drift monitoring, and live-traffic validation.")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("References", level=1)
    for ref in REFERENCES:
        numbered(doc, ref)

    doc.save(OUTPUT)
    REPORT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, REPORT_OUTPUT)


def main() -> None:
    metrics, probabilities, y_test, feature_names, train, test = run_analysis()
    make_charts(metrics, probabilities, y_test, feature_names)
    create_document(metrics, train, test)
    print(f"Created {OUTPUT}")
    print(f"Created {METRICS_CSV}")


if __name__ == "__main__":
    main()
